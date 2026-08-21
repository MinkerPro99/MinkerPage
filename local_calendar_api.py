import asyncio
import datetime as dt
import hashlib
import hmac
import io
import json
import os
import re
import secrets
import smtplib
import uuid
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path
from typing import Any

# Load .env if present
_env_path = os.path.join(os.path.dirname(__file__), ".env")
if os.path.exists(_env_path):
    with open(_env_path) as _f:
        for _line in _f:
            _line = _line.strip()
            if _line and not _line.startswith("#") and "=" in _line:
                _k, _v = _line.split("=", 1)
                os.environ.setdefault(_k.strip(), _v.strip())

import requests as http_requests
from flask import Flask, jsonify, request, send_file
from mysql.connector import Error, pooling
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename

#TEST
# DB_CONFIG = {
#     "host": os.getenv("DB_HOST", "127.0.0.1"),
#     "port": int(os.getenv("DB_PORT", "3306")),
#     "user": os.getenv("DB_USER", "root"),
#     "password": os.getenv("DB_PASSWORD", "Init.1234"),
#     "database": os.getenv("DB_NAME", "minker_calendar_prod2"),
# }

#PROD
DB_CONFIG = {
    "host": os.getenv("DB_HOST", "127.0.0.1"),
    "port": int(os.getenv("DB_PORT", "3306")),
    "user": os.getenv("DB_USER", "minker_api2"),
    "password": os.getenv("DB_PASSWORD", "Init.12345!"),
    "database": os.getenv("DB_NAME", "minker_calendar_prod2"),
}

APP_PORT = int(os.getenv("APP_PORT", "5050"))
TOKEN_DAYS = int(os.getenv("TOKEN_DAYS", "30"))
EMAIL_CODE_TTL_MINUTES = int(os.getenv("EMAIL_CODE_TTL_MINUTES", "10"))
PASSWORD_RESET_TTL_MINUTES = int(os.getenv("PASSWORD_RESET_TTL_MINUTES", "15"))
MAX_EMAIL_CODE_ATTEMPTS = int(os.getenv("MAX_EMAIL_CODE_ATTEMPTS", "5"))
EMAIL_CODE_SIGNING_SECRET = os.getenv("EMAIL_CODE_SIGNING_SECRET", "minker-local-email-secret")
STUDY_TRAINER_STORE = Path(os.getenv("STUDY_TRAINER_STORE", Path(__file__).with_name("data") / "study_trainer.json"))
STUDY_TRAINER_MAX_TEXT = int(os.getenv("STUDY_TRAINER_MAX_TEXT", "60000"))
DONE_MARKER = "\u2063\u2064\u2063"

app = Flask(__name__)

pool = pooling.MySQLConnectionPool(
    pool_name="minker_calendar_pool",
    pool_size=5,
    **DB_CONFIG,
)

def json_error(message: str, status: int = 400):
    safe_message = "Internal server error" if status >= 500 else message
    return jsonify({"ok": False, "error": safe_message}), status


def parse_iso_date(date_value: str) -> datetime.date:
    return datetime.strptime(date_value, "%Y-%m-%d").date()


def normalize_email(value: str | None) -> str:
    return (value or "").strip().lower()


def is_valid_email(value: str) -> bool:
    if not value or len(value) > 255 or any(ch.isspace() for ch in value):
        return False
    if value.count("@") != 1:
        return False
    local_part, domain = value.split("@", 1)
    if not local_part or not domain or domain.startswith(".") or domain.endswith("."):
        return False
    if "." not in domain:
        return False
    return True


def generate_six_digit_code() -> str:
    return f"{secrets.randbelow(1_000_000):06d}"


def hash_one_time_code(purpose: str, email: str, code: str) -> str:
    payload = f"{purpose}|{email}|{code}"
    return hmac.new(
        EMAIL_CODE_SIGNING_SECRET.encode("utf-8"),
        payload.encode("utf-8"),
        hashlib.sha256,
    ).hexdigest()


def codes_match(stored_hash: str, purpose: str, email: str, code: str) -> bool:
    expected_hash = hash_one_time_code(purpose, email, code)
    return hmac.compare_digest(stored_hash, expected_hash)


def send_email_message(to_email: str, subject: str, body: str) -> None:
    smtp_host = os.getenv("SMTP_HOST", "").strip()
    if not smtp_host:
        raise RuntimeError("SMTP_HOST is not configured")

    smtp_port = int(os.getenv("SMTP_PORT", "587"))
    smtp_user = os.getenv("SMTP_USER", "").strip()
    smtp_password = os.getenv("SMTP_PASSWORD", "")
    smtp_from = os.getenv("SMTP_FROM", smtp_user or "no-reply@minkerpage.local").strip()
    use_tls = os.getenv("SMTP_USE_TLS", "true").strip().lower() not in {"0", "false", "no"}

    message = EmailMessage()
    message["From"] = smtp_from
    message["To"] = to_email
    message["Subject"] = subject
    message.set_content(body)

    with smtplib.SMTP(smtp_host, smtp_port, timeout=15) as smtp:
        smtp.ehlo()
        if use_tls:
            smtp.starttls()
            smtp.ehlo()
        if smtp_user and smtp_password:
            smtp.login(smtp_user, smtp_password)
        smtp.send_message(message)


def create_and_store_email_code(
    cursor,
    *,
    user_id: int,
    email: str,
    purpose: str,
    ttl_minutes: int,
) -> str:
    cursor.execute(
        """
        UPDATE auth_email_codes
        SET used_at = UTC_TIMESTAMP()
        WHERE user_id = %s
          AND email = %s
          AND purpose = %s
          AND used_at IS NULL
        """,
        (user_id, email, purpose),
    )
    code = generate_six_digit_code()
    code_hash = hash_one_time_code(purpose, email, code)
    cursor.execute(
        """
        INSERT INTO auth_email_codes (user_id, email, purpose, code_hash, expires_at, max_attempts)
        VALUES (%s, %s, %s, %s, DATE_ADD(UTC_TIMESTAMP(), INTERVAL %s MINUTE), %s)
        """,
        (user_id, email, purpose, code_hash, ttl_minutes, MAX_EMAIL_CODE_ATTEMPTS),
    )
    return code


def has_recent_email_code_request(cursor, *, user_id: int, email: str, purpose: str) -> bool:
    cursor.execute(
        """
        SELECT COUNT(*) AS request_count
        FROM auth_email_codes
        WHERE user_id = %s
          AND email = %s
          AND purpose = %s
          AND created_at > DATE_SUB(UTC_TIMESTAMP(), INTERVAL 1 MINUTE)
        """,
        (user_id, email, purpose),
    )
    row = cursor.fetchone() or {}
    return int(row.get("request_count", 0)) > 0


def ensure_auth_schema() -> None:
    conn = None
    cursor = None
    try:
        conn = pool.get_connection()
        cursor = conn.cursor()
        cursor.execute(
            """
            SELECT COUNT(*)
            FROM information_schema.COLUMNS
            WHERE TABLE_SCHEMA = DATABASE()
              AND TABLE_NAME = 'users'
              AND COLUMN_NAME = 'email'
            """
        )
        has_email = int((cursor.fetchone() or [0])[0]) > 0
        if not has_email:
            cursor.execute("ALTER TABLE users ADD COLUMN email VARCHAR(255) NULL")

        cursor.execute(
            """
            SELECT COUNT(*)
            FROM information_schema.STATISTICS
            WHERE TABLE_SCHEMA = DATABASE()
              AND TABLE_NAME = 'users'
              AND INDEX_NAME = 'uq_users_email'
            """
        )
        has_email_index = int((cursor.fetchone() or [0])[0]) > 0
        if not has_email_index:
            cursor.execute("ALTER TABLE users ADD UNIQUE KEY uq_users_email (email)")

        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS auth_email_codes (
                code_id BIGINT UNSIGNED NOT NULL AUTO_INCREMENT,
                user_id BIGINT UNSIGNED NOT NULL,
                email VARCHAR(255) NOT NULL,
                purpose VARCHAR(40) NOT NULL,
                code_hash CHAR(64) NOT NULL,
                attempt_count INT UNSIGNED NOT NULL DEFAULT 0,
                max_attempts INT UNSIGNED NOT NULL DEFAULT 5,
                expires_at DATETIME NOT NULL,
                used_at DATETIME NULL,
                created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
                PRIMARY KEY (code_id),
                KEY idx_auth_email_codes_user_id (user_id),
                KEY idx_auth_email_codes_lookup (user_id, email, purpose, expires_at),
                CONSTRAINT fk_auth_email_codes_user
                    FOREIGN KEY (user_id) REFERENCES users(user_id)
                    ON DELETE CASCADE
            ) ENGINE=InnoDB
            """
        )

        conn.commit()
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()


ensure_auth_schema()
def parse_bearer_token() -> str | None:
    auth = request.headers.get("Authorization", "")
    if not auth.startswith("Bearer "):
        return None
    return auth.split(" ", 1)[1].strip() or None


def get_authenticated_user_id() -> int | None:
    token = parse_bearer_token()
    if not token:
        return None

    conn = None
    cursor = None
    try:
        conn = pool.get_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute(
            """
            SELECT user_id
            FROM auth_tokens
            WHERE token = %s
              AND expires_at > UTC_TIMESTAMP()
            LIMIT 1
            """,
            (token,),
        )
        row = cursor.fetchone()
        if not row:
            return None
        return int(row["user_id"])
    except Error:
        return None
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()


def require_auth_user_id() -> int:
    user_id = get_authenticated_user_id()
    if not user_id:
        raise PermissionError("Unauthorized")
    return user_id


def purge_expired_calendar_events(cursor, user_id: int) -> None:
    cursor.execute(
        """
        DELETE FROM calendar_events
        WHERE user_id = %s
          AND (
            (
              title LIKE %s
              AND end_date < DATE_SUB(CURDATE(), INTERVAL 1 MONTH)
            )
            OR
            (
              (LOWER(title) LIKE %s OR LOWER(title) LIKE %s)
              AND end_date < DATE_SUB(CURDATE(), INTERVAL 1 DAY)
            )
          )
        """,
        (
            user_id,
            f"%{DONE_MARKER}%",
            "%prüfung%",
            "%pruefung%",
        ),
    )


@app.after_request
def add_cors_headers(response):
    # Allows local browser testing from static pages while developing.
    response.headers["Access-Control-Allow-Origin"] = "*"
    response.headers["Access-Control-Allow-Headers"] = "Content-Type, Authorization"
    response.headers["Access-Control-Allow-Methods"] = "GET, POST, PUT, DELETE, OPTIONS"
    return response


@app.before_request
def handle_preflight():
    if request.method == "OPTIONS":
        return ("", 204)
    return None


@app.route("/api/health-db", methods=["GET"])
def health_db():
    try:
        conn = pool.get_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute("SELECT 1 AS ok, CURRENT_TIMESTAMP AS now_ts")
        row = cursor.fetchone()
        cursor.close()
        conn.close()
        return jsonify({"ok": True, "db": row})
    except Error as exc:
        return json_error(f"DB connection failed: {exc}", 500)


@app.route("/api/auth/register", methods=["POST"])
def register_user():
    data: dict[str, Any] = request.get_json(silent=True) or {}
    username = (data.get("username") or "").strip().lower()
    password = data.get("password") or ""

    if len(username) < 3:
        return json_error("username must be at least 3 characters")
    if len(password) < 6:
        return json_error("password must be at least 6 characters")

    conn = None
    cursor = None
    try:
        conn = pool.get_connection()
        cursor = conn.cursor(dictionary=True)

        cursor.execute("SELECT user_id FROM users WHERE username = %s LIMIT 1", (username,))
        if cursor.fetchone():
            return json_error("username already exists", 409)

        pwd_hash = generate_password_hash(password)
        cursor.execute(
            """
            INSERT INTO users (username, password_hash, email)
            VALUES (%s, %s, NULL)
            """,
            (username, pwd_hash),
        )
        conn.commit()
        return jsonify({"ok": True, "user_id": cursor.lastrowid, "username": username, "email": None}), 201
    except Error as exc:
        return json_error(f"Failed to register user: {exc}", 500)
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()


@app.route("/api/auth/login", methods=["POST"])
def login_user():
    data: dict[str, Any] = request.get_json(silent=True) or {}
    username = (data.get("username") or "").strip().lower()
    password = data.get("password") or ""

    conn = None
    cursor = None
    try:
        conn = pool.get_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute(
            """
            SELECT user_id, username, email, password_hash
            FROM users
            WHERE username = %s
            LIMIT 1
            """,
            (username,),
        )
        row = cursor.fetchone()
        if not row or not check_password_hash(row["password_hash"], password):
            return json_error("invalid username or password", 401)

        token = secrets.token_urlsafe(48)
        cursor.execute(
            """
            INSERT INTO auth_tokens (user_id, token, expires_at)
            VALUES (%s, %s, DATE_ADD(UTC_TIMESTAMP(), INTERVAL %s DAY))
            """,
            (row["user_id"], token, TOKEN_DAYS),
        )
        conn.commit()

        return jsonify(
            {
                "ok": True,
                "token": token,
                "expires_in_days": TOKEN_DAYS,
                "user": {"user_id": row["user_id"], "username": row["username"], "email": row.get("email")},
            }
        )
    except Error as exc:
        return json_error(f"Failed to login: {exc}", 500)
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()


@app.route("/api/auth/me", methods=["GET"])
def who_am_i():
    user_id = get_authenticated_user_id()
    if not user_id:
        return json_error("Unauthorized", 401)

    conn = None
    cursor = None
    try:
        conn = pool.get_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute(
            "SELECT user_id, username, email, created_at FROM users WHERE user_id = %s LIMIT 1",
            (user_id,),
        )
        row = cursor.fetchone()
        if not row:
            return json_error("User not found", 404)
        return jsonify({"ok": True, "user": row})
    except Error as exc:
        return json_error(f"Failed to fetch profile: {exc}", 500)
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()


@app.route("/api/auth/email/request-link-code", methods=["POST"])
def request_email_link_code():
    try:
        user_id = require_auth_user_id()
    except PermissionError:
        return json_error("Unauthorized", 401)

    data: dict[str, Any] = request.get_json(silent=True) or {}
    email = normalize_email(data.get("email"))
    if not is_valid_email(email):
        return json_error("Please provide a valid email address")

    conn = None
    cursor = None
    try:
        conn = pool.get_connection()
        cursor = conn.cursor(dictionary=True)

        cursor.execute("SELECT user_id FROM users WHERE email = %s LIMIT 1", (email,))
        owner = cursor.fetchone()
        if owner and int(owner["user_id"]) != user_id:
            return json_error("Email is already in use", 409)

        if has_recent_email_code_request(cursor, user_id=user_id, email=email, purpose="link_email"):
            return json_error("Please wait before requesting another code.", 429)

        code = create_and_store_email_code(
            cursor,
            user_id=user_id,
            email=email,
            purpose="link_email",
            ttl_minutes=EMAIL_CODE_TTL_MINUTES,
        )
        send_email_message(
            email,
            "Your MinkerPage verification code",
            (
                f"Use this 6-digit code to link your email to your MinkerPage account: {code}\n\n"
                f"This code expires in {EMAIL_CODE_TTL_MINUTES} minutes."
            ),
        )
        conn.commit()
        return jsonify({"ok": True, "message": "Verification code sent"})
    except RuntimeError as exc:
        if conn:
            conn.rollback()
        app.logger.warning("Email delivery unavailable: %s", exc)
        return json_error("Email delivery is unavailable. Please try again later.", 503)
    except Error as exc:
        if conn:
            conn.rollback()
        app.logger.exception("Failed to request verification code")
        return json_error(f"Failed to request verification code: {exc}", 500)
    except Exception as exc:
        if conn:
            conn.rollback()
        app.logger.exception("Failed to send verification code")
        return json_error(f"Failed to send verification code: {exc}", 500)
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()


@app.route("/api/auth/email/verify-link-code", methods=["POST"])
def verify_email_link_code():
    try:
        user_id = require_auth_user_id()
    except PermissionError:
        return json_error("Unauthorized", 401)

    data: dict[str, Any] = request.get_json(silent=True) or {}
    email = normalize_email(data.get("email"))
    code = (data.get("code") or "").strip()

    if not is_valid_email(email):
        return json_error("Please provide a valid email address")
    if not re.fullmatch(r"\d{6}", code):
        return json_error("Verification code must be 6 digits")

    conn = None
    cursor = None
    try:
        conn = pool.get_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute(
            """
            SELECT code_id, code_hash, attempt_count, max_attempts
            FROM auth_email_codes
            WHERE user_id = %s
              AND email = %s
              AND purpose = 'link_email'
              AND used_at IS NULL
              AND expires_at > UTC_TIMESTAMP()
            ORDER BY created_at DESC
            LIMIT 1
            """,
            (user_id, email),
        )
        row = cursor.fetchone()
        if not row:
            return json_error("Verification code is invalid or expired", 400)
        if int(row["attempt_count"]) >= int(row["max_attempts"]):
            return json_error("Too many invalid attempts. Request a new code.", 429)

        if not codes_match(row["code_hash"], "link_email", email, code):
            cursor.execute(
                "UPDATE auth_email_codes SET attempt_count = attempt_count + 1 WHERE code_id = %s",
                (row["code_id"],),
            )
            conn.commit()
            return json_error("Verification code is invalid", 400)

        cursor.execute("UPDATE users SET email = %s WHERE user_id = %s", (email, user_id))
        cursor.execute("UPDATE auth_email_codes SET used_at = UTC_TIMESTAMP() WHERE code_id = %s", (row["code_id"],))
        conn.commit()
        return jsonify({"ok": True, "email": email})
    except Error as exc:
        if conn:
            conn.rollback()
        if getattr(exc, "errno", None) == 1062:
            return json_error("Email is already in use", 409)
        return json_error(f"Failed to verify email: {exc}", 500)
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()


def require_email_linked(cursor, user_id: int):
    cursor.execute("SELECT username, email, password_hash FROM users WHERE user_id = %s LIMIT 1", (user_id,))
    user_row = cursor.fetchone()
    if not user_row:
        return None, json_error("User not found", 404)
    if not user_row.get("email"):
        return None, json_error("Please link an email before changing credentials", 403)
    return user_row, None


@app.route("/api/auth/settings/username", methods=["POST"])
def update_username():
    try:
        user_id = require_auth_user_id()
    except PermissionError:
        return json_error("Unauthorized", 401)

    data: dict[str, Any] = request.get_json(silent=True) or {}
    new_username = (data.get("new_username") or "").strip().lower()
    confirm_username = (data.get("confirm_username") or "").strip().lower()
    if len(new_username) < 3:
        return json_error("username must be at least 3 characters")
    if new_username != confirm_username:
        return json_error("username entries do not match")

    conn = None
    cursor = None
    try:
        conn = pool.get_connection()
        cursor = conn.cursor(dictionary=True)

        _, denied = require_email_linked(cursor, user_id)
        if denied:
            return denied

        cursor.execute(
            "SELECT user_id FROM users WHERE username = %s AND user_id <> %s LIMIT 1",
            (new_username, user_id),
        )
        if cursor.fetchone():
            return json_error("username already exists", 409)

        cursor.execute("UPDATE users SET username = %s WHERE user_id = %s", (new_username, user_id))
        conn.commit()
        return jsonify({"ok": True, "username": new_username})
    except Error as exc:
        if conn:
            conn.rollback()
        return json_error(f"Failed to update username: {exc}", 500)
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()


@app.route("/api/auth/settings/password", methods=["POST"])
def update_password():
    try:
        user_id = require_auth_user_id()
    except PermissionError:
        return json_error("Unauthorized", 401)

    data: dict[str, Any] = request.get_json(silent=True) or {}
    old_password = data.get("old_password") or ""
    new_password = data.get("new_password") or ""
    confirm_password = data.get("confirm_password") or ""
    if len(new_password) < 6:
        return json_error("new password must be at least 6 characters")
    if new_password != confirm_password:
        return json_error("new password entries do not match")

    conn = None
    cursor = None
    try:
        conn = pool.get_connection()
        cursor = conn.cursor(dictionary=True)

        user_row, denied = require_email_linked(cursor, user_id)
        if denied:
            return denied

        if not check_password_hash(user_row["password_hash"], old_password):
            return json_error("old password is incorrect", 401)

        password_hash = generate_password_hash(new_password)
        cursor.execute("UPDATE users SET password_hash = %s WHERE user_id = %s", (password_hash, user_id))

        current_token = parse_bearer_token()
        if current_token:
            cursor.execute(
                "DELETE FROM auth_tokens WHERE user_id = %s AND token <> %s",
                (user_id, current_token),
            )
        else:
            cursor.execute("DELETE FROM auth_tokens WHERE user_id = %s", (user_id,))

        conn.commit()
        return jsonify({"ok": True, "message": "Password updated"})
    except Error as exc:
        if conn:
            conn.rollback()
        return json_error(f"Failed to update password: {exc}", 500)
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()


@app.route("/api/auth/password/forgot", methods=["POST"])
def forgot_password():
    data: dict[str, Any] = request.get_json(silent=True) or {}
    email = normalize_email(data.get("email"))
    generic_response = {
        "ok": True,
        "message": "If this email exists, a reset code has been sent.",
    }

    if not is_valid_email(email):
        return jsonify(generic_response)

    conn = None
    cursor = None
    try:
        conn = pool.get_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute("SELECT user_id FROM users WHERE email = %s LIMIT 1", (email,))
        user_row = cursor.fetchone()
        if not user_row:
            return jsonify(generic_response)

        user_id = int(user_row["user_id"])
        if has_recent_email_code_request(cursor, user_id=user_id, email=email, purpose="reset_password"):
            return jsonify(generic_response)

        code = create_and_store_email_code(
            cursor,
            user_id=user_id,
            email=email,
            purpose="reset_password",
            ttl_minutes=PASSWORD_RESET_TTL_MINUTES,
        )
        send_email_message(
            email,
            "Your MinkerPage password reset code",
            (
                f"Use this 6-digit code to reset your MinkerPage password: {code}\n\n"
                f"This code expires in {PASSWORD_RESET_TTL_MINUTES} minutes."
            ),
        )
        conn.commit()
        return jsonify(generic_response)
    except RuntimeError as exc:
        if conn:
            conn.rollback()
        app.logger.warning("Password reset email delivery unavailable: %s", exc)
        return json_error("Email delivery is unavailable. Please try again later.", 503)
    except Error as exc:
        if conn:
            conn.rollback()
        app.logger.exception("Failed to process password reset request")
        return json_error(f"Failed to process password reset request: {exc}", 500)
    except Exception as exc:
        if conn:
            conn.rollback()
        app.logger.exception("Failed to send password reset email")
        return json_error(f"Failed to send password reset email: {exc}", 500)
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()


@app.route("/api/auth/password/reset", methods=["POST"])
def reset_password_with_code():
    data: dict[str, Any] = request.get_json(silent=True) or {}
    email = normalize_email(data.get("email"))
    code = (data.get("code") or "").strip()
    new_password = data.get("new_password") or ""
    confirm_password = data.get("confirm_password") or ""

    if not is_valid_email(email):
        return json_error("Please provide a valid email address")
    if not re.fullmatch(r"\d{6}", code):
        return json_error("Reset code must be 6 digits")
    if len(new_password) < 6:
        return json_error("password must be at least 6 characters")
    if new_password != confirm_password:
        return json_error("password entries do not match")

    conn = None
    cursor = None
    try:
        conn = pool.get_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute("SELECT user_id FROM users WHERE email = %s LIMIT 1", (email,))
        user_row = cursor.fetchone()
        if not user_row:
            return json_error("Invalid reset code or expired request", 400)
        user_id = int(user_row["user_id"])

        cursor.execute(
            """
            SELECT code_id, code_hash, attempt_count, max_attempts
            FROM auth_email_codes
            WHERE user_id = %s
              AND email = %s
              AND purpose = 'reset_password'
              AND used_at IS NULL
              AND expires_at > UTC_TIMESTAMP()
            ORDER BY created_at DESC
            LIMIT 1
            """,
            (user_id, email),
        )
        code_row = cursor.fetchone()
        if not code_row:
            return json_error("Invalid reset code or expired request", 400)
        if int(code_row["attempt_count"]) >= int(code_row["max_attempts"]):
            return json_error("Too many invalid attempts. Request a new code.", 429)

        if not codes_match(code_row["code_hash"], "reset_password", email, code):
            cursor.execute(
                "UPDATE auth_email_codes SET attempt_count = attempt_count + 1 WHERE code_id = %s",
                (code_row["code_id"],),
            )
            conn.commit()
            return json_error("Invalid reset code", 400)

        cursor.execute(
            "UPDATE users SET password_hash = %s WHERE user_id = %s",
            (generate_password_hash(new_password), user_id),
        )
        cursor.execute("DELETE FROM auth_tokens WHERE user_id = %s", (user_id,))
        cursor.execute("UPDATE auth_email_codes SET used_at = UTC_TIMESTAMP() WHERE code_id = %s", (code_row["code_id"],))
        conn.commit()
        return jsonify({"ok": True, "message": "Password reset successful. Please log in."})
    except Error as exc:
        if conn:
            conn.rollback()
        return json_error(f"Failed to reset password: {exc}", 500)
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()


@app.route("/api/events", methods=["GET"])
def list_events():
    try:
        user_id = require_auth_user_id()
    except PermissionError:
        return json_error("Unauthorized", 401)

    start = request.args.get("start")
    end = request.args.get("end")

    try:
        conn = pool.get_connection()
        cursor = conn.cursor(dictionary=True)
        purge_expired_calendar_events(cursor, user_id)
        conn.commit()

        if start and end:
            parse_iso_date(start)
            parse_iso_date(end)
            sql = """
                SELECT event_id, title, start_date, end_date, description, color_hex, created_at, updated_at
                FROM calendar_events
                WHERE user_id = %s
                  AND start_date <= %s
                  AND end_date >= %s
                ORDER BY start_date, event_id
            """
            cursor.execute(sql, (user_id, end, start))
        else:
            sql = """
                SELECT event_id, title, start_date, end_date, description, color_hex, created_at, updated_at
                FROM calendar_events
                WHERE user_id = %s
                ORDER BY start_date, event_id
            """
            cursor.execute(sql, (user_id,))

        rows = cursor.fetchall()
        cursor.close()
        conn.close()
        return jsonify({"ok": True, "events": rows})
    except ValueError:
        return json_error("Invalid date format. Use YYYY-MM-DD.")
    except Error as exc:
        return json_error(f"Failed to fetch events: {exc}", 500)


@app.route("/api/events", methods=["POST"])
def create_event():
    try:
        user_id = require_auth_user_id()
    except PermissionError:
        return json_error("Unauthorized", 401)

    data: dict[str, Any] = request.get_json(silent=True) or {}
    title = (data.get("title") or "").strip()
    start_date = data.get("start_date")
    end_date = data.get("end_date")
    description = data.get("description")
    color_hex = data.get("color_hex")

    if not title:
        return json_error("title is required")
    if not start_date:
        return json_error("start_date is required")
    if not end_date:
        end_date = start_date

    try:
        s_date = parse_iso_date(start_date)
        e_date = parse_iso_date(end_date)
        if e_date < s_date:
            return json_error("end_date must be >= start_date")

        conn = pool.get_connection()
        cursor = conn.cursor()
        cursor.execute(
            """
            INSERT INTO calendar_events (user_id, title, start_date, end_date, description, color_hex)
            VALUES (%s, %s, %s, %s, %s, %s)
            """,
            (user_id, title, start_date, end_date, description, color_hex),
        )
        conn.commit()
        new_id = cursor.lastrowid
        cursor.close()
        conn.close()
        return jsonify({"ok": True, "event_id": new_id}), 201
    except ValueError:
        return json_error("Invalid date format. Use YYYY-MM-DD.")
    except Error as exc:
        return json_error(f"Failed to create event: {exc}", 500)


@app.route("/api/events/<int:event_id>", methods=["PUT"])
def update_event(event_id: int):
    try:
        user_id = require_auth_user_id()
    except PermissionError:
        return json_error("Unauthorized", 401)

    data: dict[str, Any] = request.get_json(silent=True) or {}
    title = (data.get("title") or "").strip()
    start_date = data.get("start_date")
    end_date = data.get("end_date")
    description = data.get("description")
    color_hex = data.get("color_hex")

    if not title:
        return json_error("title is required")
    if not start_date:
        return json_error("start_date is required")
    if not end_date:
        end_date = start_date

    try:
        s_date = parse_iso_date(start_date)
        e_date = parse_iso_date(end_date)
        if e_date < s_date:
            return json_error("end_date must be >= start_date")

        conn = pool.get_connection()
        cursor = conn.cursor()
        cursor.execute(
            """
            UPDATE calendar_events
            SET title = %s, start_date = %s, end_date = %s, description = %s, color_hex = %s
            WHERE event_id = %s AND user_id = %s
            """,
            (title, start_date, end_date, description, color_hex, event_id, user_id),
        )
        conn.commit()
        updated_rows = cursor.rowcount
        cursor.close()
        conn.close()

        if updated_rows == 0:
            return json_error("Event not found", 404)

        return jsonify({"ok": True, "updated": True})
    except ValueError:
        return json_error("Invalid date format. Use YYYY-MM-DD.")
    except Error as exc:
        return json_error(f"Failed to update event: {exc}", 500)


@app.route("/api/events/<int:event_id>", methods=["DELETE"])
def delete_event(event_id: int):
    try:
        user_id = require_auth_user_id()
    except PermissionError:
        return json_error("Unauthorized", 401)

    try:
        conn = pool.get_connection()
        cursor = conn.cursor()
        cursor.execute(
            "DELETE FROM calendar_events WHERE event_id = %s AND user_id = %s",
            (event_id, user_id),
        )
        conn.commit()
        deleted_rows = cursor.rowcount
        cursor.close()
        conn.close()

        if deleted_rows == 0:
            return json_error("Event not found", 404)

        return jsonify({"ok": True, "deleted": True})
    except Error as exc:
        return json_error(f"Failed to delete event: {exc}", 500)


# ── Study Trainer ─────────────────────────────────────────────────────────────

def _utc_now_iso() -> str:
    return dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat()


def _load_study_trainer_store() -> dict[str, Any]:
    if not STUDY_TRAINER_STORE.exists():
        return {"version": 1, "users": {}}

    try:
        with STUDY_TRAINER_STORE.open("r", encoding="utf-8") as f:
            data = json.load(f)
    except (json.JSONDecodeError, OSError):
        return {"version": 1, "users": {}}

    if not isinstance(data, dict):
        return {"version": 1, "users": {}}
    data.setdefault("version", 1)
    data.setdefault("users", {})
    return data


def _save_study_trainer_store(data: dict[str, Any]) -> None:
    STUDY_TRAINER_STORE.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = STUDY_TRAINER_STORE.with_suffix(".tmp")
    with tmp_path.open("w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    tmp_path.replace(STUDY_TRAINER_STORE)


def _get_user_study_data(store: dict[str, Any], user_id: int) -> dict[str, Any]:
    users = store.setdefault("users", {})
    key = str(user_id)
    if key not in users:
        users[key] = {"subjects": [], "sessions": [], "updated_at": _utc_now_iso()}
    users[key].setdefault("subjects", [])
    users[key].setdefault("sessions", [])
    return users[key]


def _find_subject(user_data: dict[str, Any], subject_id: str) -> dict[str, Any] | None:
    return next((subject for subject in user_data.get("subjects", []) if subject.get("id") == subject_id), None)


def _normalise_subject_exams(subject: dict[str, Any]) -> dict[str, Any]:
    exams = subject.setdefault("exams", {})
    if not isinstance(exams, dict):
        exams = {}
        subject["exams"] = exams

    for event_id in subject.get("event_ids", []) or []:
        key = str(event_id)
        exam = exams.setdefault(key, {})
        exam.setdefault("event_id", int(event_id))
        exam.setdefault("mastery", None)
        exam.setdefault("links", [])
        exam.setdefault("prompt", "")
        exam.setdefault("chat_messages", [])
        exam.setdefault("notes", "")
        exam.setdefault("files", [])
        exam.setdefault("file_text", "")
        exam.setdefault("mock_exam", None)
        exam.setdefault("answers", [])
        exam.setdefault("insights", {})
        exam.setdefault("created_at", _utc_now_iso())
        exam.setdefault("updated_at", _utc_now_iso())

    return exams


def _get_subject_exam(subject: dict[str, Any], event_id: int) -> dict[str, Any]:
    exams = _normalise_subject_exams(subject)
    key = str(event_id)
    exam = exams.setdefault(
        key,
        {
            "event_id": event_id,
            "mastery": None,
            "links": [],
            "prompt": "",
            "chat_messages": [],
            "notes": "",
            "files": [],
            "file_text": "",
            "mock_exam": None,
            "answers": [],
            "insights": {},
            "created_at": _utc_now_iso(),
            "updated_at": _utc_now_iso(),
        },
    )
    if event_id not in [int(value) for value in subject.get("event_ids", []) or []]:
        subject["event_ids"] = sorted(set([*subject.get("event_ids", []), event_id]))
    return exam


def _subject_average_mastery(subject: dict[str, Any]) -> int:
    exams = _normalise_subject_exams(subject)
    scores = [
        int(exam.get("mastery"))
        for exam in exams.values()
        if isinstance(exam, dict) and exam.get("mastery") is not None
    ]
    if not scores:
        return max(0, min(100, int(subject.get("mastery", 35) or 35)))
    return round(sum(scores) / len(scores))


def _fetch_calendar_events_for_user(user_id: int, event_ids: list[int] | None = None) -> list[dict[str, Any]]:
    conn = None
    cursor = None
    try:
        conn = pool.get_connection()
        cursor = conn.cursor(dictionary=True)
        purge_expired_calendar_events(cursor, user_id)
        conn.commit()

        if event_ids:
            placeholders = ", ".join(["%s"] * len(event_ids))
            cursor.execute(
                f"""
                SELECT event_id, title, start_date, end_date, description, color_hex
                FROM calendar_events
                WHERE user_id = %s
                  AND event_id IN ({placeholders})
                ORDER BY start_date, event_id
                """,
                (user_id, *event_ids),
            )
        else:
            today = dt.date.today().isoformat()
            cursor.execute(
                """
                SELECT event_id, title, start_date, end_date, description, color_hex
                FROM calendar_events
                WHERE user_id = %s
                  AND start_date >= %s
                ORDER BY start_date, event_id
                LIMIT 40
                """,
                (user_id, today),
            )

        return cursor.fetchall()
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()


def _event_to_json(event: dict[str, Any]) -> dict[str, Any]:
    def clean_date(value: Any) -> str:
        if hasattr(value, "isoformat"):
            return value.isoformat()
        return str(value or "")

    return {
        "event_id": int(event["event_id"]),
        "title": event.get("title") or "Untitled event",
        "start_date": clean_date(event.get("start_date")),
        "end_date": clean_date(event.get("end_date")),
        "description": event.get("description") or "",
        "color_hex": event.get("color_hex") or "",
    }


def _extract_uploaded_file_text() -> tuple[list[dict[str, str]], str]:
    file_summaries: list[dict[str, str]] = []
    collected: list[str] = []
    readable_exts = {".txt", ".md", ".csv", ".json", ".html", ".htm", ".log"}

    for file in request.files.getlist("files"):
        filename = secure_filename(file.filename or "uploaded-file")
        raw = file.read()
        ext = Path(filename).suffix.lower()

        if ext in readable_exts:
            text = raw.decode("utf-8", errors="ignore")
        elif ext == ".pdf":
            try:
                from pypdf import PdfReader

                reader = PdfReader(io.BytesIO(raw))
                text = "\n".join(page.extract_text() or "" for page in reader.pages)
            except Exception:
                text = f"[{filename}: PDF uploaded, but text extraction is unavailable.]"
        elif ext == ".docx":
            try:
                import docx

                document = docx.Document(io.BytesIO(raw))
                text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            except Exception:
                text = f"[{filename}: Word document uploaded, but text extraction is unavailable.]"
        else:
            text = f"[{filename}: binary file uploaded. The filename may still indicate the topic.]"

        text = text[:20000]
        file_summaries.append({"filename": filename, "characters_used": str(len(text))})
        collected.append(f"File: {filename}\n{text}")

    return file_summaries, "\n\n".join(collected)


def _compact_source_text(*parts: str) -> str:
    text = "\n\n".join(part.strip() for part in parts if part and part.strip())
    text = re.sub(r"\s+", " ", text).strip()
    return text[:STUDY_TRAINER_MAX_TEXT]


def _fetch_link_source_text(links: list[str]) -> str:
    collected: list[str] = []
    for link in links[:6]:
        if not re.match(r"^https?://", link, re.I):
            collected.append(f"Link: {link}\n[Skipped: only http and https links can be read.]")
            continue
        try:
            response = http_requests.get(
                link,
                timeout=8,
                headers={"User-Agent": "MinkerPage StudyTrainer/1.0"},
            )
            response.raise_for_status()
            content_type = response.headers.get("content-type", "")
            raw = response.text[:120000]
            if "html" in content_type.lower():
                raw = re.sub(r"(?is)<(script|style|noscript).*?</\1>", " ", raw)
                raw = re.sub(r"(?is)<[^>]+>", " ", raw)
            raw = re.sub(r"\s+", " ", raw).strip()
            collected.append(f"Link: {link}\n{raw[:14000]}")
        except Exception as exc:
            collected.append(f"Link: {link}\n[Could not read linked source: {exc}]")
    return "\n\n".join(collected)


def _extract_source_topics(source_text: str, limit: int = 8) -> list[str]:
    stop_words = {
        "about", "after", "also", "because", "before", "between", "could", "from", "have", "into",
        "more", "most", "only", "should", "that", "their", "there", "these", "this", "with",
        "your", "will", "would", "exam", "study", "material", "question", "answer",
    }
    words = re.findall(r"[A-Za-zÄÖÜäöüß][A-Za-zÄÖÜäöüß\-]{3,}", source_text.lower())
    counts: dict[str, int] = {}
    for word in words:
        clean = word.strip("-")
        if clean in stop_words:
            continue
        counts[clean] = counts.get(clean, 0) + 1
    return [word for word, _count in sorted(counts.items(), key=lambda item: item[1], reverse=True)[:limit]]


def _fallback_questions(source_text: str, count: int = 8) -> list[dict[str, Any]]:
    sentences = [
        s.strip()
        for s in re.split(r"(?<=[.!?])\s+", source_text)
        if len(s.strip()) > 40
    ]
    topics = _extract_source_topics(source_text, count)
    if not sentences:
        sentences = [
            "The available source material is short. Build answers from the uploaded files, links, and prompt.",
            "Use the central definitions, formulas, examples, and relationships from the source collection.",
            "Focus on applying the material instead of repeating isolated facts.",
        ]
    if not topics:
        topics = ["core concept", "definitions", "applications", "connections", "common mistakes"]

    stems = [
        "Explain the concept of {topic} and show how it would appear in an exam task.",
        "Compare {topic} with a related idea from the sources. What is the key difference?",
        "Solve a realistic exam-style problem involving {topic}. Explain each step.",
        "Identify a common mistake about {topic} and correct it with evidence from the sources.",
        "Create a short example that demonstrates {topic}, then explain why it works.",
        "Summarize the most important rule or definition for {topic} and apply it.",
        "Connect {topic} to another source topic and explain the relationship.",
        "Answer as if this were a timed exam: what are the required points for {topic}?",
    ]
    questions = []
    for index in range(count):
        topic = topics[index % len(topics)]
        evidence = sentences[index % len(sentences)][:260].rstrip()
        questions.append(
            {
                "id": f"q{index + 1}",
                "topic": topic.title(),
                "difficulty": "medium" if index < 5 else "hard",
                "prompt": stems[index % len(stems)].format(topic=topic),
                "expected_answer": evidence,
                "hint": "Use the sources to build an answer with definitions, reasoning, and an example where possible.",
            }
        )
    return questions


def _fallback_plan(events: list[dict[str, Any]], source_text: str) -> dict[str, Any]:
    event_titles = ", ".join((event.get("title") or "selected exam") for event in events) or "selected exam"
    questions = _fallback_questions(source_text)
    return {
        "summary": f"Study plan for {event_titles}.",
        "study_plan": [
            {"day": 1, "focus": "Map the material", "tasks": ["Skim all sources", "Write a topic checklist", "Answer questions 1-2"]},
            {"day": 2, "focus": "Practice recall", "tasks": ["Create flashcards", "Answer questions 3-5", "Review weak answers"]},
            {"day": 3, "focus": "Exam simulation", "tasks": ["Answer questions 6-8 without notes", "Revise the weakest topic"]},
        ],
        "questions": questions,
        "insights": {
            "strengths": [],
            "needs_work": ["No AI provider was configured, so this plan was generated locally from the supplied text."],
            "next_tasks": ["Add more source text or configure GROQ_API_KEY for richer question generation."],
        },
    }


def _normalise_ai_plan(plan: dict[str, Any], events: list[dict[str, Any]], source_text: str) -> dict[str, Any]:
    fallback = _fallback_plan(events, source_text)
    if not isinstance(plan, dict):
        return fallback

    questions = plan.get("questions")
    if not isinstance(questions, list) or not questions:
        questions = fallback["questions"]

    normalised_questions = []
    for index, question in enumerate(questions[:12]):
        if not isinstance(question, dict):
            continue
        normalised_questions.append(
            {
                "id": str(question.get("id") or f"q{index + 1}"),
                "topic": str(question.get("topic") or "Core material"),
                "difficulty": str(question.get("difficulty") or "medium"),
                "prompt": str(question.get("prompt") or question.get("question") or fallback["questions"][index % len(fallback["questions"])]["prompt"]),
                "expected_answer": str(question.get("expected_answer") or question.get("answer") or ""),
                "hint": str(question.get("hint") or "Review the source material and explain the reasoning step by step."),
            }
        )

    return {
        "summary": str(plan.get("summary") or fallback["summary"]),
        "study_plan": plan.get("study_plan") if isinstance(plan.get("study_plan"), list) else fallback["study_plan"],
        "questions": normalised_questions or fallback["questions"],
        "insights": plan.get("insights") if isinstance(plan.get("insights"), dict) else fallback["insights"],
    }


def _generate_ai_study_plan(events: list[dict[str, Any]], source_text: str, links: list[str], notes: str) -> dict[str, Any]:
    groq_key = os.getenv("GROQ_API_KEY", "")
    if not groq_key:
        return _fallback_plan(events, source_text)

    prompt = f"""You are an advanced exam study assistant similar to NotebookLM, but stricter about testing understanding.
Analyze the student's prompt, uploaded files, and linked source text. Create a mock exam that SYNTHESIZES the material.
Do not copy the user's prompt as a question. Do not ask generic questions like "explain the material".
Questions must be realistic exam questions with a clear topic, rubric, expected answer, and hint.

Return strict JSON only with this schema:
{{
  "summary": "short summary",
  "study_plan": [{{"day": 1, "focus": "topic", "tasks": ["task"]}}],
  "questions": [{{"id": "q1", "topic": "topic", "difficulty": "easy|medium|hard", "prompt": "question", "expected_answer": "answer rubric", "hint": "hint"}}],
  "insights": {{"strengths": [], "needs_work": [], "next_tasks": []}}
}}
Selected calendar events:
{json.dumps(events, ensure_ascii=False)}
Links:
{json.dumps(links, ensure_ascii=False)}
User notes:
{notes}
Source material:
{source_text}
"""

    try:
        resp = http_requests.post(
            "https://api.groq.com/openai/v1/chat/completions",
            headers={"Authorization": f"Bearer {groq_key}", "Content-Type": "application/json"},
            json={
                "model": os.getenv("STUDY_TRAINER_MODEL", "llama-3.1-8b-instant"),
                "messages": [{"role": "user", "content": prompt}],
                "temperature": 0.35,
                "max_tokens": 2400,
                "response_format": {"type": "json_object"},
            },
            timeout=35,
        )
        resp.raise_for_status()
        raw = resp.json()["choices"][0]["message"]["content"].strip()
        return _normalise_ai_plan(json.loads(raw), events, source_text)
    except Exception as exc:
        print(f"[StudyTrainer] AI generation failed: {exc}")
        return _fallback_plan(events, source_text)


def _review_answers_with_ai(questions: list[dict[str, Any]], submitted: list[dict[str, str]], source_text: str) -> list[dict[str, Any]] | None:
    groq_key = os.getenv("GROQ_API_KEY", "")
    if not groq_key:
        return None

    prompt = f"""You are an exam answer reviewer. Grade each student answer against the expected answer and source material.
Be strict but fair. Decide whether the answer is correct enough for exam readiness.
Return strict JSON only:
{{"results": [{{"question_id": "q1", "score": 0.0, "is_correct": false, "feedback": "short feedback", "review": "specific review", "target_points": "missing required points", "needed_area": "topic to study"}}]}}

Questions and rubrics:
{json.dumps(questions, ensure_ascii=False)}

Student answers:
{json.dumps(submitted, ensure_ascii=False)}

Source material excerpt:
{source_text[:20000]}
"""
    try:
        resp = http_requests.post(
            "https://api.groq.com/openai/v1/chat/completions",
            headers={"Authorization": f"Bearer {groq_key}", "Content-Type": "application/json"},
            json={
                "model": os.getenv("STUDY_TRAINER_REVIEW_MODEL", os.getenv("STUDY_TRAINER_MODEL", "llama-3.1-8b-instant")),
                "messages": [{"role": "user", "content": prompt}],
                "temperature": 0.1,
                "max_tokens": 2200,
                "response_format": {"type": "json_object"},
            },
            timeout=45,
        )
        resp.raise_for_status()
        raw = resp.json()["choices"][0]["message"]["content"]
        parsed = json.loads(raw)
        results = parsed.get("results")
        return results if isinstance(results, list) else None
    except Exception:
        return None


def _score_answer(user_answer: str, expected: str) -> tuple[float, str]:
    user_words = set(re.findall(r"[a-zA-Z0-9äöüÄÖÜß]{4,}", user_answer.lower()))
    expected_words = set(re.findall(r"[a-zA-Z0-9äöüÄÖÜß]{4,}", expected.lower()))
    if not user_answer.strip():
        return 0.0, "No answer submitted."
    if not expected_words:
        return (0.65 if len(user_answer.strip()) > 80 else 0.35), "Answer saved; no detailed rubric was available."

    overlap = len(user_words & expected_words) / max(1, len(expected_words))
    score = max(0.0, min(1.0, overlap))
    if score >= 0.65:
        feedback = "Good coverage of the expected key ideas."
    elif score >= 0.35:
        feedback = "Partly correct. Add more of the key terms and explain the reasoning."
    else:
        feedback = "Needs more work. Revisit the source and answer with the central terms."
    return round(score, 2), feedback


@app.route("/api/study-trainer/sessions", methods=["GET"])
def study_trainer_sessions():
    try:
        user_id = require_auth_user_id()
    except PermissionError:
        return json_error("Unauthorized", 401)

    store = _load_study_trainer_store()
    user_data = _get_user_study_data(store, user_id)
    return jsonify({"ok": True, "sessions": user_data["sessions"]})


@app.route("/api/study-trainer/subjects", methods=["GET"])
def study_trainer_subjects():
    try:
        user_id = require_auth_user_id()
    except PermissionError:
        return json_error("Unauthorized", 401)

    store = _load_study_trainer_store()
    user_data = _get_user_study_data(store, user_id)
    return jsonify({"ok": True, "subjects": user_data["subjects"]})


@app.route("/api/study-trainer/subjects", methods=["POST"])
def study_trainer_create_subject():
    try:
        user_id = require_auth_user_id()
    except PermissionError:
        return json_error("Unauthorized", 401)

    data: dict[str, Any] = request.get_json(silent=True) or {}
    name = (data.get("name") or "").strip()
    if len(name) < 2:
        return json_error("Subject name must be at least 2 characters")

    store = _load_study_trainer_store()
    user_data = _get_user_study_data(store, user_id)
    subject = {
        "id": str(uuid.uuid4()),
        "name": name,
        "event_ids": [],
        "mastery": 35,
        "exams": {},
        "created_at": _utc_now_iso(),
        "updated_at": _utc_now_iso(),
    }
    user_data["subjects"].append(subject)
    user_data["updated_at"] = _utc_now_iso()
    _save_study_trainer_store(store)
    return jsonify({"ok": True, "subject": subject}), 201


@app.route("/api/study-trainer/subjects/<subject_id>", methods=["PUT"])
def study_trainer_update_subject(subject_id: str):
    try:
        user_id = require_auth_user_id()
    except PermissionError:
        return json_error("Unauthorized", 401)

    data: dict[str, Any] = request.get_json(silent=True) or {}
    store = _load_study_trainer_store()
    user_data = _get_user_study_data(store, user_id)
    subject = _find_subject(user_data, subject_id)
    if not subject:
        return json_error("Subject not found", 404)

    if "name" in data:
        name = (data.get("name") or "").strip()
        if len(name) < 2:
            return json_error("Subject name must be at least 2 characters")
        subject["name"] = name

    if "event_ids" in data:
        event_ids = []
        for value in data.get("event_ids") or []:
            try:
                event_ids.append(int(value))
            except (TypeError, ValueError):
                continue
        subject["event_ids"] = sorted(set(event_ids))
        _normalise_subject_exams(subject)

    if "mastery" in data:
        try:
            subject["mastery"] = max(0, min(100, int(data.get("mastery"))))
        except (TypeError, ValueError):
            return json_error("mastery must be a number between 0 and 100")

    subject["updated_at"] = _utc_now_iso()
    user_data["updated_at"] = _utc_now_iso()
    _save_study_trainer_store(store)
    return jsonify({"ok": True, "subject": subject})


@app.route("/api/study-trainer/subjects/<subject_id>/complete", methods=["POST"])
def study_trainer_complete_subject(subject_id: str):
    try:
        user_id = require_auth_user_id()
    except PermissionError:
        return json_error("Unauthorized", 401)

    store = _load_study_trainer_store()
    user_data = _get_user_study_data(store, user_id)
    subject = _find_subject(user_data, subject_id)
    if not subject:
        return json_error("Subject not found", 404)

    user_data["subjects"] = [item for item in user_data["subjects"] if item.get("id") != subject_id]
    user_data["sessions"] = [item for item in user_data["sessions"] if item.get("subject_id") != subject_id]
    user_data["updated_at"] = _utc_now_iso()
    _save_study_trainer_store(store)
    return jsonify({"ok": True, "deleted_subject": subject_id})


@app.route("/api/study-trainer/subjects/<subject_id>/exams/<int:event_id>", methods=["DELETE"])
def study_trainer_remove_exam(subject_id: str, event_id: int):
    try:
        user_id = require_auth_user_id()
    except PermissionError:
        return json_error("Unauthorized", 401)

    store = _load_study_trainer_store()
    user_data = _get_user_study_data(store, user_id)
    subject = _find_subject(user_data, subject_id)
    if not subject:
        return json_error("Subject not found", 404)

    subject["event_ids"] = [int(value) for value in subject.get("event_ids", []) if int(value) != event_id]
    exams = subject.setdefault("exams", {})
    if isinstance(exams, dict):
        exams.pop(str(event_id), None)
    subject["mastery"] = _subject_average_mastery(subject)
    subject["updated_at"] = _utc_now_iso()
    user_data["updated_at"] = _utc_now_iso()
    _save_study_trainer_store(store)
    return jsonify({"ok": True, "subject": subject})


@app.route("/api/study-trainer/subjects/<subject_id>/exams/<int:event_id>/material", methods=["POST"])
def study_trainer_save_exam_material(subject_id: str, event_id: int):
    try:
        user_id = require_auth_user_id()
    except PermissionError:
        return json_error("Unauthorized", 401)

    store = _load_study_trainer_store()
    user_data = _get_user_study_data(store, user_id)
    subject = _find_subject(user_data, subject_id)
    if not subject:
        return json_error("Subject not found", 404)

    events = _fetch_calendar_events_for_user(user_id, [event_id])
    if not events:
        return json_error("Matching calendar event not found", 404)

    try:
        links = [str(link).strip() for link in json.loads(request.form.get("links", "[]")) if str(link).strip()]
    except json.JSONDecodeError:
        links = []

    prompt = (request.form.get("prompt") or "").strip()
    chat_message = (request.form.get("chat_message") or "").strip()
    notes = (request.form.get("notes") or "").strip()
    uploaded_files, file_text = _extract_uploaded_file_text()

    exam = _get_subject_exam(subject, event_id)
    exam["links"] = links
    exam["prompt"] = prompt
    exam["notes"] = notes
    if chat_message:
        messages = exam.setdefault("chat_messages", [])
        messages.append({
            "role": "user",
            "content": chat_message,
            "created_at": _utc_now_iso(),
        })
        exam["chat_messages"] = messages[-20:]
    if uploaded_files:
        exam["files"] = [*exam.get("files", []), *uploaded_files]
    if file_text:
        exam["file_text"] = _compact_source_text(exam.get("file_text", ""), file_text)
    exam["event"] = _event_to_json(events[0])
    exam["updated_at"] = _utc_now_iso()
    subject["updated_at"] = _utc_now_iso()
    user_data["updated_at"] = _utc_now_iso()
    _save_study_trainer_store(store)
    return jsonify({"ok": True, "subject": subject, "exam": exam})


@app.route("/api/study-trainer/subjects/<subject_id>/exams/<int:event_id>/generate", methods=["POST"])
def study_trainer_generate_exam(subject_id: str, event_id: int):
    try:
        user_id = require_auth_user_id()
    except PermissionError:
        return json_error("Unauthorized", 401)

    store = _load_study_trainer_store()
    user_data = _get_user_study_data(store, user_id)
    subject = _find_subject(user_data, subject_id)
    if not subject:
        return json_error("Subject not found", 404)

    events = [_event_to_json(event) for event in _fetch_calendar_events_for_user(user_id, [event_id])]
    if not events:
        return json_error("Matching calendar event not found", 404)

    exam = _get_subject_exam(subject, event_id)
    links = [str(link) for link in exam.get("links", []) if str(link).strip()]
    prompt = str(exam.get("prompt") or "")
    chat_text = "\n".join(str(message.get("content", "")) for message in exam.get("chat_messages", []) if isinstance(message, dict))
    notes = str(exam.get("notes") or "")
    file_text = str(exam.get("file_text") or "")
    link_text = _fetch_link_source_text(links)
    source_text = _compact_source_text(prompt, chat_text, notes, "\n".join(links), link_text, file_text)
    if not source_text:
        return json_error("Add files, links, or an extra prompt before generating a mock exam")

    ai_plan = _generate_ai_study_plan(events, source_text, links, prompt or notes)
    exam["event"] = events[0]
    exam["mock_exam"] = {
        "id": str(uuid.uuid4()),
        "created_at": _utc_now_iso(),
        "summary": ai_plan["summary"],
        "study_plan": ai_plan["study_plan"],
        "questions": ai_plan["questions"],
    }
    exam["answers"] = []
    exam["insights"] = ai_plan["insights"]
    exam["updated_at"] = _utc_now_iso()
    subject["updated_at"] = _utc_now_iso()
    user_data["updated_at"] = _utc_now_iso()
    _save_study_trainer_store(store)
    return jsonify({"ok": True, "subject": subject, "exam": exam})


@app.route("/api/study-trainer/subjects/<subject_id>/exams/<int:event_id>/answers", methods=["POST"])
def study_trainer_submit_exam_answers(subject_id: str, event_id: int):
    try:
        user_id = require_auth_user_id()
    except PermissionError:
        return json_error("Unauthorized", 401)

    data: dict[str, Any] = request.get_json(silent=True) or {}
    submitted = data.get("answers") or []
    if not isinstance(submitted, list):
        return json_error("answers must be an array")

    store = _load_study_trainer_store()
    user_data = _get_user_study_data(store, user_id)
    subject = _find_subject(user_data, subject_id)
    if not subject:
        return json_error("Subject not found", 404)

    exam = _get_subject_exam(subject, event_id)
    mock_exam = exam.get("mock_exam") or {}
    questions = {question["id"]: question for question in mock_exam.get("questions", [])}
    if not questions:
        return json_error("Generate a mock exam before submitting answers")

    results = []
    strengths = []
    needs_work = []
    study_material = []
    link_text = _fetch_link_source_text([str(link) for link in exam.get("links", []) if str(link).strip()])
    source_text = _compact_source_text(
        str(exam.get("prompt") or ""),
        "\n".join(str(message.get("content", "")) for message in exam.get("chat_messages", []) if isinstance(message, dict)),
        str(exam.get("notes") or ""),
        link_text,
        str(exam.get("file_text") or ""),
    )
    ai_results = _review_answers_with_ai(list(questions.values()), submitted, source_text)
    ai_by_id = {
        str(item.get("question_id") or ""): item
        for item in ai_results or []
        if isinstance(item, dict)
    }

    for item in submitted:
        question_id = str(item.get("question_id") or "")
        answer = str(item.get("answer") or "")
        question = questions.get(question_id)
        if not question:
            continue
        ai_review = ai_by_id.get(question_id)
        if ai_review:
            try:
                score = max(0.0, min(1.0, float(ai_review.get("score", 0))))
            except (TypeError, ValueError):
                score = 0.0
            feedback = str(ai_review.get("feedback") or "")
        else:
            score, feedback = _score_answer(answer, question.get("expected_answer", ""))
        topic = question.get("topic", "Core material")
        expected = question.get("expected_answer", "")
        missing_focus = expected[:220] if score < 0.65 and expected else question.get("hint", "")
        result = {
            "question_id": question_id,
            "answer": answer,
            "score": score,
            "feedback": feedback or ("Good coverage of the rubric." if score >= 0.65 else "Important rubric points are missing."),
            "is_correct": bool(ai_review.get("is_correct")) if ai_review else score >= 0.65,
            "review": str(ai_review.get("review") or "") if ai_review else ("Correct enough for exam readiness." if score >= 0.65 else "This answer misses important expected ideas. Rework it using the target points below."),
            "target_points": str(ai_review.get("target_points") or missing_focus) if ai_review else missing_focus,
            "needed_area": str(ai_review.get("needed_area") or (topic if score < 0.65 else "")) if ai_review else (topic if score < 0.65 else ""),
            "topic": topic,
            "submitted_at": _utc_now_iso(),
        }
        results.append(result)
        if score >= 0.65:
            strengths.append(topic)
        else:
            needs_work.append(topic)
            study_material.append(f"Review {topic}: {question.get('expected_answer') or question.get('hint') or 'revisit the source material.'}")

    mastery = round((sum(float(item["score"]) for item in results) / max(1, len(results))) * 100)
    exam["answers"] = results
    exam["mastery"] = mastery
    exam["insights"] = {
        "strengths": sorted(set(strengths)),
        "needs_work": sorted(set(needs_work)),
        "next_tasks": [
            f"Redo the missed questions about {topic}, then explain the concept without notes."
            for topic in sorted(set(needs_work))[:4]
        ] or ["Run one timed recap before the exam."],
        "study_material": study_material[:8],
    }
    exam["updated_at"] = _utc_now_iso()
    subject["mastery"] = _subject_average_mastery(subject)
    subject["updated_at"] = _utc_now_iso()
    user_data["updated_at"] = _utc_now_iso()
    _save_study_trainer_store(store)
    return jsonify({"ok": True, "subject": subject, "exam": exam})


@app.route("/api/study-trainer/generate", methods=["POST"])
def study_trainer_generate():
    try:
        user_id = require_auth_user_id()
    except PermissionError:
        return json_error("Unauthorized", 401)

    try:
        selected_raw = request.form.get("selected_event_ids", "[]")
        selected_event_ids = [int(value) for value in json.loads(selected_raw) if str(value).isdigit()]
    except (TypeError, ValueError, json.JSONDecodeError):
        return json_error("selected_event_ids must be a JSON array of event ids")

    if not selected_event_ids:
        return json_error("Select at least one calendar event")

    subject_id = (request.form.get("subject_id") or "").strip()

    try:
        links = [str(link).strip() for link in json.loads(request.form.get("links", "[]")) if str(link).strip()]
    except json.JSONDecodeError:
        links = []

    notes = (request.form.get("notes") or "").strip()
    events = [_event_to_json(event) for event in _fetch_calendar_events_for_user(user_id, selected_event_ids)]
    if not events:
        return json_error("No matching calendar events found", 404)

    uploaded_files, file_text = _extract_uploaded_file_text()
    source_text = _compact_source_text(notes, "\n".join(links), file_text)
    if not source_text:
        return json_error("Add notes, links, or at least one file before generating a plan")

    store = _load_study_trainer_store()
    user_data = _get_user_study_data(store, user_id)
    subject = _find_subject(user_data, subject_id) if subject_id else None
    if subject_id and not subject:
        return json_error("Selected subject was not found", 404)

    ai_plan = _generate_ai_study_plan(events, source_text, links, notes)
    session = {
        "id": str(uuid.uuid4()),
        "created_at": _utc_now_iso(),
        "updated_at": _utc_now_iso(),
        "subject_id": subject_id,
        "subject_name": "",
        "events": events,
        "links": links,
        "notes": notes,
        "files": uploaded_files,
        "summary": ai_plan["summary"],
        "study_plan": ai_plan["study_plan"],
        "questions": ai_plan["questions"],
        "insights": ai_plan["insights"],
        "answers": [],
    }

    if subject:
        session["subject_name"] = subject.get("name") or ""
        subject["event_ids"] = sorted(set([*subject.get("event_ids", []), *selected_event_ids]))
        subject["updated_at"] = _utc_now_iso()

    user_data["sessions"].insert(0, session)
    user_data["sessions"] = user_data["sessions"][:20]
    user_data["updated_at"] = _utc_now_iso()
    _save_study_trainer_store(store)

    return jsonify({"ok": True, "session": session}), 201


@app.route("/api/study-trainer/sessions/<session_id>/answers", methods=["POST"])
def study_trainer_submit_answers(session_id: str):
    try:
        user_id = require_auth_user_id()
    except PermissionError:
        return json_error("Unauthorized", 401)

    data: dict[str, Any] = request.get_json(silent=True) or {}
    submitted = data.get("answers") or []
    if not isinstance(submitted, list):
        return json_error("answers must be an array")

    store = _load_study_trainer_store()
    user_data = _get_user_study_data(store, user_id)
    session = next((item for item in user_data["sessions"] if item.get("id") == session_id), None)
    if not session:
        return json_error("Study session not found", 404)

    questions = {question["id"]: question for question in session.get("questions", [])}
    results = []
    strengths = []
    needs_work = []

    for item in submitted:
        question_id = str(item.get("question_id") or "")
        answer = str(item.get("answer") or "")
        question = questions.get(question_id)
        if not question:
            continue
        score, feedback = _score_answer(answer, question.get("expected_answer", ""))
        result = {
            "question_id": question_id,
            "answer": answer,
            "score": score,
            "feedback": feedback,
            "submitted_at": _utc_now_iso(),
        }
        results.append(result)
        if score >= 0.65:
            strengths.append(question.get("topic", "Core material"))
        else:
            needs_work.append(question.get("topic", "Core material"))

    session["answers"] = results
    session["updated_at"] = _utc_now_iso()
    session["insights"] = {
        "strengths": sorted(set(strengths)),
        "needs_work": sorted(set(needs_work)),
        "next_tasks": [
            f"Redo questions about {topic} and write a one paragraph explanation."
            for topic in sorted(set(needs_work))[:4]
        ] or ["Move to a timed exam simulation."],
    }
    user_data["updated_at"] = _utc_now_iso()
    _save_study_trainer_store(store)

    return jsonify({"ok": True, "results": results, "insights": session["insights"], "session": session})


# ── Jarvis TTS preview ────────────────────────────────────────────────────────

JARVIS_VOICE   = "en-US-ChristopherNeural"
JARVIS_RATE    = "-10%"
JARVIS_PITCH   = "-5Hz"
WEATHER_CITY   = "Mettmenstetten"
GTA6_RELEASE   = dt.date(2026, 11, 19)

MODULE_NAMES = {
    "M": "Math", "F": "French", "D": "German",
    "GS": "History", "Phys": "Physics", "E": "English",
}

def _jarvis_ordinal(n: int) -> str:
    if 11 <= n <= 13: return f"{n}th"
    return f"{n}{['th','st','nd','rd','th'][min(n % 10, 4)]}"

def _jarvis_format_date(d: dt.date) -> str:
    return d.strftime("%A, %B ") + _jarvis_ordinal(d.day)

def _jarvis_expand_title(title: str) -> str:
    title = re.sub(r'\bM(\d+)\b', lambda m: f"Module {m.group(1)}", title)
    title = re.sub(
        r'\b(' + '|'.join(re.escape(k) for k in sorted(MODULE_NAMES, key=len, reverse=True)) + r')\b(?= -| –|$)',
        lambda m: MODULE_NAMES[m.group(1)], title
    )
    for de, en in {"Prüfung": "Exam", "Aufgaben": "Tasks", "erstellen": "create",
                   "abgeben": "submit", "Kapitel": "Chapter", "Rotes buch": "Red book", "und": "and"}.items():
        title = re.sub(re.escape(de), en, title, flags=re.IGNORECASE)
    return title

def _jarvis_weather() -> str:
    try:
        r = http_requests.get(
            f"https://wttr.in/{WEATHER_CITY}?format=%C,+%t,+feels+like+%f",
            headers={"User-Agent": "curl/7.0"}, timeout=8
        )
        raw = r.text.strip()
        raw = re.sub(r'\+(-?\d+)°C', lambda m: f"{m.group(1)} degrees", raw)
        raw = re.sub(r'(-\d+)°C', lambda m: f"minus {m.group(1)[1:]} degrees", raw)
        return raw
    except Exception:
        return ""

def _jarvis_calendar(user_id: int) -> str:
    try:
        today = dt.date.today()
        end   = today + dt.timedelta(days=7)
        conn   = pool.get_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            SELECT title, start_date, description FROM calendar_events
            WHERE user_id = %s AND start_date >= %s AND start_date <= %s
            ORDER BY start_date, event_id
        """, (user_id, today.isoformat(), end.isoformat()))
        rows = cursor.fetchall()
        cursor.close(); conn.close()
        if not rows:
            return "No upcoming events in the next 7 days."
        lines = []
        for row in rows:
            title = _jarvis_expand_title(row["title"])
            d     = row["start_date"]
            date  = d.date() if hasattr(d, 'date') else dt.date.fromisoformat(str(d)[:10])
            lines.append(f"{_jarvis_format_date(date)}: {title}")
        return "\n".join(f"- {l}" for l in lines)
    except Exception as e:
        return ""

def _jarvis_build_script(user_id: int) -> str:
    now      = dt.datetime.now()
    hour     = now.hour
    tod      = "morning" if hour < 12 else "afternoon" if hour < 17 else "evening" if hour < 21 else "night"
    day_str  = _jarvis_format_date(now.date()) + now.strftime(" at %I:%M %p").replace(" 0", " ").lstrip()
    days     = (GTA6_RELEASE - now.date()).days
    gta6     = f"Grand Theft Auto 6 is releasing in {days} days." if days > 0 else "Grand Theft Auto 6 has already been released."
    weather  = _jarvis_weather()
    calendar = _jarvis_calendar(user_id)

    groq_key = os.getenv("GROQ_API_KEY", "")
    if groq_key:
        try:
            news = _jarvis_fetch_news()
            news_section     = f"Current gaming headlines: {news}" if news else ""
            weather_section  = f"Today's weather in {WEATHER_CITY}: {weather}" if weather else ""
            calendar_section = f"Upcoming study calendar events (next 7 days):\n{calendar}"
            prompt = f"""You are JARVIS, Tony Stark's AI. Write a single short spoken greeting (4-5 sentences, no more).
            Today is {day_str}. Use the exact day names provided for calendar events — do not infer "tomorrow" or "next week" yourself.
            Context:
            - Current date and time: {day_str}
            - {weather_section}
            - {calendar_section}
            - {news_section}
            - GTA 6 aside: {gta6}
            Rules:
                - Start with "Good {tod} sir."
                - say that all systems are online, (somewhat like the voice in the game subnautica).
                - Naturally mention the date and time early on.
                - Briefly mention the weather in one clause.
                - Read out ALL calendar events listed, grouped by day, using the exact day names given. Do not skip any.
                - Pick ONE gaming headline and slip it in briefly.
                - You MUST include the GTA 6 fact as a short aside — this is mandatory, do not skip it.
                - End with a short offer of assistance as a statement, not a question (e.g., "I'm here if you need me").
                - Tone: Sophisticated, understated, and impeccably polite, with a vein of dry, British-style wit. No asterisks, no markdown, plain text only."""

            resp = http_requests.post(
                "https://api.groq.com/openai/v1/chat/completions",
                headers={"Authorization": f"Bearer {groq_key}", "Content-Type": "application/json"},
                json={"model": "llama-3.1-8b-instant", "messages": [{"role": "user", "content": prompt}], "max_tokens": 400},
                timeout=15,
            )
            resp.raise_for_status()
            text = resp.json()["choices"][0]["message"]["content"].strip()
            if text:
                return text
        except Exception as e:
            print(f"[Jarvis] Groq failed: {e}")

    # Fallback template
    import random
    openers = [
        f"Good {tod}, sir. It is {day_str}.",
        f"Good {tod}, sir. I've been expecting you. It is {day_str}.",
    ]
    parts = [random.choice(openers)]
    if weather:
        parts.append(f"Weather in {WEATHER_CITY}: {weather}.")
    parts.append(f"Upcoming schedule: {calendar}")
    parts.append(gta6)
    parts.append("All systems online.")
    return "  ".join(parts)


def _jarvis_fetch_news(n: int = 3) -> str:
    try:
        import xml.etree.ElementTree as ET
        r     = http_requests.get("https://www.pcgamer.com/rss/", headers={"User-Agent": "Mozilla/5.0"}, timeout=8)
        root  = ET.fromstring(r.content)
        items = root.findall(".//item")[:n]
        return " | ".join(i.findtext("title", "").strip() for i in items if i.findtext("title"))
    except Exception:
        return ""

async def _jarvis_synthesise(text: str) -> bytes:
    import edge_tts
    buf = io.BytesIO()
    communicate = edge_tts.Communicate(text, JARVIS_VOICE, rate=JARVIS_RATE, pitch=JARVIS_PITCH)
    async for chunk in communicate.stream():
        if chunk["type"] == "audio":
            buf.write(chunk["data"])
    buf.seek(0)
    return buf.read()

@app.route("/api/jarvis/greet", methods=["GET"])
def jarvis_greet():
    user_id = get_authenticated_user_id()
    if not user_id:
        return json_error("Unauthorized", 401)
    try:
        script = _jarvis_build_script(user_id)
        audio  = asyncio.run(_jarvis_synthesise(script))
        return send_file(io.BytesIO(audio), mimetype="audio/mpeg", as_attachment=False)
    except Exception as e:
        return json_error(f"Jarvis failed: {e}", 500)


@app.route("/api/jarvis/script", methods=["GET"])
def jarvis_script():
    user_id = get_authenticated_user_id()
    if not user_id:
        return json_error("Unauthorized", 401)
    try:
        return jsonify({"ok": True, "script": _jarvis_build_script(user_id)})
    except Exception as e:
        return json_error(f"Jarvis failed: {e}", 500)


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=APP_PORT, debug=True)
